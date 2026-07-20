import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
BASE="/sessions/gifted-happy-euler/mnt/kylas-crm-mcp-server/presales_june_analysis"
rows=json.load(open(BASE+"/open_pipeline_focus.json"))
NAVY="1F3864"; BLUE="2E5496"; TEAL="1F6E6E"; GREEN="C6EFCE"; GREENF="006100"
AMB="FFEB9C"; AMBF="9C6500"; RED="FFC7CE"; REDF="9C0006"; WHITE="FFFFFF"; GOLD="FFF2CC"; GREY="F2F2F2"
thin=Side(style="thin",color="BFBFBF"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def hf(s=10,c=WHITE,b=True): return Font(name="Arial",size=s,bold=b,color=c)
def f(s=9,c="000000",b=False): return Font(name="Arial",size=s,bold=b,color=c)
wrap=Alignment(wrap_text=True,vertical="top"); ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center")
ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=Workbook()
bandfill={"HOT":GREEN,"WARM":AMB,"MID":GREY,"COLD":RED}
bandfont={"HOT":GREENF,"WARM":AMBF,"MID":"000000","COLD":REDF}
churnfill={"KEEP":GREEN,"REASSIGN-NEW-SALES":AMB,"BACK-TO-PRESALES":"D9E1F2","DISQUALIFY":RED}
def title(ws,t,span,c=NAVY):
    x=ws.cell(row=1,column=1,value=t); x.font=hf(13,WHITE,True); x.fill=PatternFill("solid",fgColor=c)
    x.alignment=left; ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=span); ws.row_dimensions[1].height=24

def table(ws,recs,cols,startrow=3,colorband=True,colorchurn=True):
    for i,(h,w,_key) in enumerate(cols):
        ws.column_dimensions[get_column_letter(i+1)].width=w
        x=ws.cell(row=startrow,column=i+1,value=h); x.font=hf(9,WHITE,True); x.fill=PatternFill("solid",fgColor=BLUE); x.alignment=ctrw; x.border=border
    ws.row_dimensions[startrow].height=26
    r=startrow+1
    for rec in recs:
        for i,key in enumerate([c[2] for c in cols]):
            v=rec.get(key)
            if key=="conv_prob" and v is not None: v=v
            x=ws.cell(row=r,column=i+1,value=v); x.border=border; x.font=f(9)
            x.alignment=wrap if cols[i][1]>=34 else (left if key in("name","next_action","churn_reason","one_line","owner","city","stage") else ctr)
            if key=="conv_prob" and v is not None: x.number_format="0%"
            if key=="band" and colorband and v in bandfill: x.fill=PatternFill("solid",fgColor=bandfill[v]); x.font=f(9,bandfont[v],True)
            if key=="churn_call" and colorchurn and v in churnfill: x.fill=PatternFill("solid",fgColor=churnfill[v]); x.font=f(9, REDF if v=="DISQUALIFY" else ("006100" if v=="KEEP" else "000000"),True)
            if key=="verdict" and v in bandfill: pass
        r+=1
    ws.freeze_panes=ws.cell(row=startrow+1,column=1)
    return ws

# 1. Summary
ws=wb.active; ws.title="Focus Summary"; ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[30,14,16,60]): ws.column_dimensions[c].width=w
title(ws,"Open Pre-Sales Pipeline — Triage & Focus (150 qualified, active, 5+ seat, demo-done deals)",4)
import collections
band=collections.Counter(r["band"] for r in rows); churn=collections.Counter(r["churn_call"] for r in rows)
val=lambda sub: sum((r["licences"] or 0)*1700 for r in sub)
def hrow(r,label,v,note,fillc=None,fc="000000"):
    ws.cell(row=r,column=1,value=label).font=f(10,NAVY,True); ws.cell(row=r,column=1).border=border; ws.cell(row=r,column=1).alignment=left
    c2=ws.cell(row=r,column=2,value=v); c2.border=border; c2.alignment=ctr; c2.font=f(10,fc,True)
    if fillc: c2.fill=PatternFill("solid",fgColor=fillc)
    ws.cell(row=r,column=4,value=note).font=f(10); ws.cell(row=r,column=4).alignment=wrap; ws.cell(row=r,column=4).border=border
    ws.cell(row=r,column=3).border=border
r=3
ws.cell(row=r,column=1,value="Total open deals in scope").font=f(10,NAVY,True); ws.cell(row=r,column=2,value=150).font=f(10,b=True)
ws.cell(row=r,column=4,value="Open, forecast=OPEN, demo conducted, 5+ licences, touched since 15-May. (The full open book is 480; this is the workable slice.)").alignment=wrap; r+=2
ws.cell(row=r,column=1,value="CONVERSION BANDS").font=hf(10,WHITE,True); ws.cell(row=r,column=1).fill=PatternFill("solid",fgColor=TEAL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); r+=1
for b in ["HOT","WARM","MID","COLD"]:
    sub=[x for x in rows if x["band"]==b]; hrow(r,b,len(sub),f"~INR {val(sub):,} ACV",bandfill[b],bandfont[b]); r+=1
r+=1
ws.cell(row=r,column=1,value="CHURN CALL").font=hf(10,WHITE,True); ws.cell(row=r,column=1).fill=PatternFill("solid",fgColor=TEAL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); r+=1
cnote={"KEEP":"Keep with current sales — execute next action","REASSIGN-NEW-SALES":"Qualified but stalled on the rep — hand to a sharper closer","BACK-TO-PRESALES":"Never properly qualified (BANT missing / wrong contact) — re-qualify","DISQUALIFY":"Dead / sub-ICP / non-responsive — close-lost, stop spending demos"}
for cc in ["KEEP","REASSIGN-NEW-SALES","BACK-TO-PRESALES","DISQUALIFY"]:
    sub=[x for x in rows if x["churn_call"]==cc]; hrow(r,cc,len(sub),cnote[cc]+f"  (~INR {val(sub):,})",churnfill[cc]); r+=1
r+=1
wn=[x for x in rows if x["band"] in("HOT","WARM")]
ws.cell(row=r,column=1,value="WORK NOW (HOT+WARM)").font=f(10,NAVY,True); ws.cell(row=r,column=2,value=len(wn)).font=f(10,GREENF,True); ws.cell(row=r,column=2).fill=PatternFill("solid",fgColor=GOLD)
ws.cell(row=r,column=4,value=f"~INR {val(wn):,} ACV — the deals sales should chase this week. Only ~21% of the 'qualified' book is genuinely workable now; ~49% is mis-qualified, stalled or dead.").alignment=wrap; ws.row_dimensions[r].height=44

cols_full=[("Prob",7,"conv_prob"),("Band",7,"band"),("Deal ID",9,"deal_id"),("Name",20,"name"),("Seats",6,"licences"),("City",13,"city"),("Type",11,"dev_broker"),("Stage",15,"stage"),("Owner",15,"owner"),("BANT",8,"bant_flag"),("Idle",6,"idle_days"),("Verdict",8,"verdict"),("Churn Call",17,"churn_call"),("Next action (Sales)",52,"next_action"),("Churn reason",50,"churn_reason")]

# 2. WORK NOW
ws=wb.create_sheet("WORK NOW"); ws.sheet_view.showGridLines=False
title(ws,"WORK NOW — HOT + WARM (chase this week), ranked by probability",len(cols_full),TEAL)
table(ws,[r for r in rows if r["band"] in("HOT","WARM")],cols_full)
# 3. All 150
ws=wb.create_sheet("All 150 (ranked)"); ws.sheet_view.showGridLines=False
title(ws,"All 150 open deals — ranked by conversion probability",len(cols_full))
table(ws,rows,cols_full)
# 4/5/6 churn buckets
for name,cc,c in [("Reassign to New Sales","REASSIGN-NEW-SALES",AMBF),("Back to Pre-Sales","BACK-TO-PRESALES",BLUE),("Disqualify","DISQUALIFY",REDF)]:
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    sub=[r for r in rows if r["churn_call"]==cc]
    title(ws,f"{name} — {len(sub)} deals",len(cols_full),c)
    table(ws,sorted(sub,key=lambda x:-(x['conv_prob'] or 0)),cols_full)
wb.save(BASE+"/Open_Pipeline_Focus.xlsx")
print("saved:",wb.sheetnames,"| rows",len(rows))
