import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
BASE="/sessions/gifted-happy-euler/mnt/kylas-crm-mcp-server/presales_june_analysis"
rows=json.load(open(BASE+"/lost_dataset.json")); A=json.load(open(BASE+"/lost_analytics.json"))
NAVY="1F3864";BLUE="2E5496";TEAL="1F6E6E";GREENBG="E2EFDA";AMB="FFF2CC";REDBG="FBE4E4";WHITE="FFFFFF";GREENF="1E7145";REDF="9C0006";GOLD="FFF4D6"
thin=Side(style="thin",color="BFBFBF");bd=Border(left=thin,right=thin,top=thin,bottom=thin)
def hf(s=10,c=WHITE,b=True):return Font(name="Arial",size=s,bold=b,color=c)
def f(s=9,c="000000",b=False):return Font(name="Arial",size=s,bold=b,color=c)
wrap=Alignment(wrap_text=True,vertical="top");ctr=Alignment(horizontal="center",vertical="center");left=Alignment(horizontal="left",vertical="center");ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=Workbook()
def title(ws,t,span,c=NAVY):
    x=ws.cell(row=1,column=1,value=t);x.font=hf(13,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=left
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=span);ws.row_dimensions[1].height=24
def band(ws,t,span,row,c=BLUE):
    x=ws.cell(row=row,column=1,value=t);x.font=hf(11,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=left
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=span);ws.row_dimensions[row].height=20
def head(ws,hs,row,c=BLUE):
    for i,h in enumerate(hs):
        x=ws.cell(row=row,column=i+1,value=h);x.font=hf(9,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=ctrw;x.border=bd
def cv(ws,r,c,v,b=False,al=None,fill=None,fc="000000",nf=None):
    x=ws.cell(row=r,column=c,value=v);x.border=bd;x.font=f(9,fc,b);x.alignment=al or ctr
    if fill:x.fill=PatternFill("solid",fgColor=fill)
    if nf:x.number_format=nf

# Sheet 1 — Summary
ws=wb.active;ws.title="Summary";ws.sheet_view.showGridLines=False
for c,w in zip("ABCDE",[34,13,13,13,42]):ws.column_dimensions[c].width=w
title(ws,"Lost / Unqualified Pre-Sales Deals — lifecycle & sales effort (5+ seats, created since 1 Apr 2026)",5)
band(ws,"Scope",5,2)
cv(ws,3,1,"Deals in scope",True,left);cv(ws,3,2,A["N"]);ws.cell(row=3,column=2).font=f(10,b=True)
cv(ws,3,5,"Pre-sales team, 5+ licences, created Apr–Jun 2026, now Closed Lost or Closed Unqualified.",al=wrap)
cv(ws,4,1,"Closed Lost / Closed Unqualified",True,left);cv(ws,4,2,f'{A["lost"]} / {A["unq"]}')
cv(ws,4,5,"52% Lost (worked then lost), 48% Unqualified (killed as not-a-fit).",al=wrap)
r=6
band(ws,"Lifecycle timing (days) — created → demo → lost",5,r);head(ws,["Stage","Created→Demo (med)","Demo→Lost (med)","Total cycle (med)","No demo held"],r+1);r+=2
tm=A["timing"]
for lab,key,extra in [("All",'all',A["no_demo"]),("Closed Lost",'lost',None),("Closed Unqualified",'unq',None)]:
    t=tm[key]
    cv(ws,r,1,lab,True,left);cv(ws,r,2,f'{t["cd"][1]}d (avg {t["cd"][0]})');cv(ws,r,3,f'{t["dl"][1]}d (avg {t["dl"][0]})');cv(ws,r,4,f'{t["tt"][1]}d (avg {t["tt"][0]})')
    cv(ws,r,5,(f'{extra} of {A["N"]} had NO demo' if extra is not None else ""))
    r+=1
cv(ws,r,1,"Note: created→demo is exact (from fields). Lost date uses the deal's last-update timestamp as a proxy (the connector doesn't expose actualClosureDate), so demo→lost & total are close estimates.",al=wrap);ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=5);ws.cell(row=r,column=1).font=f(8,"666666",False);ws.row_dimensions[r].height=28;r+=2
band(ws,"Sales follow-up before giving up",5,r);head(ws,["Metric","Value","","",""],r+1);r+=2
sm=[("Avg attempts after demo",f'{A["attempts_avg"]}  (median {A["attempts_median"]}, max 24)',GOLD),
    ("Given up after ≤2 attempts","137 of 282 (49%)",REDBG),
    ("Zero post-demo attempts","47 (17%)",REDBG),
    ("Avg attempts — Closed Lost","4.7",None),
    ("Avg attempts — Closed Unqualified","3.0",None)]
for m,v,fl in sm:
    cv(ws,r,1,m,True,left);cv(ws,r,2,v,True,fill=fl);cv(ws,r,3,"");cv(ws,r,4,"");cv(ws,r,5,"");r+=1

# Sheet 2 — Loss reasons + effort
ws=wb.create_sheet("Loss reasons & effort");ws.sheet_view.showGridLines=False
for c,w in zip("ABCDE",[28,10,10,16,16]):ws.column_dimensions[c].width=w
title(ws,"Why deals were lost — and how hard sales chased each",5,TEAL)
head(ws,["Loss reason","Deals","% ","Avg sales attempts","Demo→lost (median)"],2)
order=sorted(A["loss_reason"].items(),key=lambda kv:-kv[1])
rr=3
for k,v in order:
    cv(ws,rr,1,k,True,left);cv(ws,rr,2,v);cv(ws,rr,3,v/A["N"],nf="0%");cv(ws,rr,4,A["attempts_by_reason"].get(k,""))
    # median demo->lost per reason from dataset
    ds=[r["d_demo_lost"] for r in rows if r["loss_reason"]==k and r["d_demo_lost"] is not None and r["d_demo_lost"]>=0]
    import statistics as st
    cv(ws,rr,5,(int(st.median(ds)) if ds else ""))
    if k=="Non-responsive":
        for cc in range(1,6):ws.cell(row=rr,column=cc).fill=PatternFill("solid",fgColor=AMB)
    rr+=1
rr+=1
band(ws,"Attempts distribution (all 282)",5,rr);head(ws,["Attempts bucket","Deals","%","",""],rr+1);rr+=2
for lab,key in [("0 (no follow-up)","0"),("1–2","1-2"),("3–5","3-5"),("6–10","6-10"),("11+","11+")]:
    n=A["attempts_dist"][key];cv(ws,rr,1,lab,True,left);cv(ws,rr,2,n);cv(ws,rr,3,n/A["N"],nf="0%");cv(ws,rr,4,"");cv(ws,rr,5,"")
    if key=="0":ws.cell(row=rr,column=1).fill=PatternFill("solid",fgColor=REDBG)
    rr+=1

# Sheet 3 — All deals
ws=wb.create_sheet("All 282 deals");ws.freeze_panes="A2"
cols=[("Deal ID",10),("Name",20),("Sales Owner",16),("Stage",12),("Seats",6),("City",13),("Create→Demo",11),("Demo→Lost",11),("Total days",10),("Sales attempts",12),("Loss reason",20),("One-line",60)]
for i,(h,w) in enumerate(cols):
    ws.column_dimensions[get_column_letter(i+1)].width=w
    x=ws.cell(row=1,column=i+1,value=h);x.font=hf(9,WHITE,True);x.fill=PatternFill("solid",fgColor=BLUE);x.alignment=ctrw;x.border=bd
ws.row_dimensions[1].height=26
rr=2
for r in sorted(rows,key=lambda x:(x["stage"],-(x["attempts"] or 0))):
    vals=[r["deal_id"],r["name"],r["owner"],r["stage"],r["licences"],r["city"],
          (int(r["d_create_demo"]) if r["d_create_demo"] is not None and r["d_create_demo"]>=0 else ("no demo" if not r["has_demo"] else "")),
          (int(r["d_demo_lost"]) if r["d_demo_lost"] is not None and r["d_demo_lost"]>=0 else ""),
          (int(r["d_total"]) if r["d_total"] is not None and r["d_total"]>=0 else ""),
          r["attempts"],r["loss_reason"],r["one_line"]]
    for ci,v in enumerate(vals,1):
        x=ws.cell(row=rr,column=ci,value=v);x.border=bd;x.font=f(9);x.alignment=wrap if ci in(2,12) else (left if ci in(3,6,11) else ctr)
    sc=ws.cell(row=rr,column=4)
    if r["stage"]=="Lost":sc.fill=PatternFill("solid",fgColor=REDBG);sc.font=f(9,REDF,True)
    else:sc.fill=PatternFill("solid",fgColor=AMB)
    if (r["attempts"] or 0)==0:ws.cell(row=rr,column=10).font=f(9,REDF,True)
    rr+=1
ws.auto_filter.ref=f"A1:L{rr-1}"
wb.save(BASE+"/Lost_Deals_Lifecycle.xlsx")
print("saved",wb.sheetnames)
