import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
BASE="/sessions/gifted-happy-euler/mnt/kylas-crm-mcp-server/marketing_weekly"
WK="2026-07-12"
rows=json.load(open(f"{BASE}/deals_{WK}.json")); S=json.load(open(f"{BASE}/snapshot_{WK}.json"))
NAVY="1F3864";BLUE="2E5496";TEAL="1F6E6E";GREENBG="E2EFDA";AMB="FFF2CC";REDBG="FBE4E4";WHITE="FFFFFF";GREENF="1E7145";REDF="9C0006";GOLD="FFF4D6"
thin=Side(style="thin",color="BFBFBF");bd=Border(left=thin,right=thin,top=thin,bottom=thin)
def hf(s=10,c=WHITE,b=True):return Font(name="Arial",size=s,bold=b,color=c)
def f(s=9,c="000000",b=False):return Font(name="Arial",size=s,bold=b,color=c)
wrap=Alignment(wrap_text=True,vertical="top");ctr=Alignment(horizontal="center",vertical="center");left=Alignment(horizontal="left",vertical="center");ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=Workbook()
def title(ws,t,span,c=NAVY):
    x=ws.cell(row=1,column=1,value=t);x.font=hf(12,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=left
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=span);ws.row_dimensions[1].height=24
def head(ws,hs,row,c=BLUE):
    for i,h in enumerate(hs):
        x=ws.cell(row=row,column=i+1,value=h);x.font=hf(9,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=ctrw;x.border=bd
def cv(ws,r,c,v,b=False,al=None,fill=None,fc="000000",nf=None):
    x=ws.cell(row=r,column=c,value=v);x.border=bd;x.font=f(9,fc,b);x.alignment=al or ctr
    if fill:x.fill=PatternFill("solid",fgColor=fill)
    if nf:x.number_format=nf
N=S["N"];T=S["targets"]
# 1 Scorecard
ws=wb.active;ws.title="Scorecard";ws.sheet_view.showGridLines=False
for c,w in zip("ABCDE",[32,12,12,12,52]):ws.column_dimensions[c].width=w
title(ws,f"Marketing & Pre-Sales — week {S['week_start']} to {S['week_end']}",5)
head(ws,["Metric","Actual","Target","Status","Read"],2)
def sc(r,m,a,t,ok,read):
    cv(ws,r,1,m,True,left);cv(ws,r,2,a,True,fill=(GREENBG if ok else REDBG),fc=(GREENF if ok else REDF))
    cv(ws,r,3,t);cv(ws,r,4,"ON" if ok else "MISS",True,fc=(GREENF if ok else REDF));cv(ws,r,5,read,al=wrap)
r=3
sc(r,"Demos conducted",N,"—",True,"43 demos — down vs the ~59/wk June run-rate");r+=1
sc(r,"Demos of 5+ licences",S["five"],T["five"],S["five"]>=T["five"],"52% of target — the headline miss");r+=1
sc(r,"GREEN (well-qualified)",S["green"],T["green"],S["green"]>=T["green"],"Only 16% of demos were GREEN (June: 39%)");r+=1
sc(r,"RED (should not be booked)",S["verdict"].get("RED",0),"<10%",False,"37% RED — materially worse than June's 22%");r+=1
sc(r,"Developer share",f'{S["dev"]}/{N} ({S["dev"]/N:.0%})',"≥50%",S["dev"]/N>=T["devshare"],"42% — below the floor; CP-heavy again");r+=1
sc(r,"Onsite (Main-6)",f'{S["onsite_m6"]}/{S["main6"]} ({S["onsite_m6"]/max(1,S["main6"]):.0%})',"≥40%",False,"15% — the on-ground model still isn't happening");r+=1
sc(r,"HOT + WARM (worth chasing)",S["band"].get("HOT",0)+S["band"].get("WARM",0),"—",True,"7 deals; 84% of the week landed COLD");r+=1
r+=1
cv(ws,r,1,"ALERTS",True,left,fill=REDBG);cv(ws,r,2,"",);cv(ws,r,3,"");cv(ws,r,4,"");cv(ws,r,5,"");r+=1
al=S["alerts"]
for k,lab,rd in [("zero_followup","Zero sales follow-up since demo","3 of them are HOT/WARM — act today"),
                 ("unverified","Contact unverified / phone mismatch","28% of the week — qualification leak"),
                 ("bad_name","Junk / single-word contact names","21% of the week"),
                 ("dummy_phone","Dummy phone (+919999999999)","Should be auto-rejected at booking")]:
    cv(ws,r,1,lab,True,left);cv(ws,r,2,al[k],True,fill=REDBG,fc=REDF);cv(ws,r,3,"0");cv(ws,r,4,"MISS",True,fc=REDF);cv(ws,r,5,rd,al=wrap);r+=1

# 2 Source & campaign
ws=wb.create_sheet("Source");ws.sheet_view.showGridLines=False
for c,w in zip("ABCDEFGH",[14,9,9,9,9,9,11,44]):ws.column_dimensions[c].width=w
title(ws,"Channel performance — last week",8,TEAL)
head(ws,["Source","Demos","5+","GREEN","RED","Dev","HOT/WARM","Read"],2)
notes={"Organic":"Best channel again — most 5+, most developers, most HOT/WARM. Protect & grow.",
 "Google":"Best size mix (7 of 10 are 5+). Reliable paid lever — scale.",
 "Facebook":"11 demos → 1 GREEN, 5 RED, 1 developer, ZERO HOT/WARM. Confirms June (2.1% win). Cut or re-target.",
 "WhatsApp":"Only 1 of 7 is 5+ and 4 are RED — inbound WA is bringing tiny/unqualified leads.",
 "Unknown":"Source not set — fix tagging."}
r=3
for s,v in sorted(S["source"].items(),key=lambda kv:-kv[1]["demos"]):
    fill=REDBG if s in("Facebook","WhatsApp") else (GREENBG if s in("Organic","Google") else None)
    cv(ws,r,1,s,True,left,fill);cv(ws,r,2,v["demos"]);cv(ws,r,3,v["five"]);cv(ws,r,4,v["green"]);cv(ws,r,5,v["red"]);cv(ws,r,6,v["dev"]);cv(ws,r,7,v["hotwarm"])
    cv(ws,r,8,notes.get(s,""),al=wrap);r+=1

# 3 Allocation & routing
ws=wb.create_sheet("Allocation & Routing");ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[26,12,12,54]):ws.column_dimensions[c].width=w
title(ws,"Allocation by sales person, and lead routing",4,TEAL)
head(ws,["Sales owner","Demos","Even-split target","Note"],2)
tgt=round(N/len(S["alloc"]),1); r=3
for o,c in sorted(S["alloc"].items(),key=lambda kv:-kv[1]):
    cv(ws,r,1,o,True,left);cv(ws,r,2,c);cv(ws,r,3,tgt)
    cv(ws,r,4,("Under-fed" if c<tgt-1 else ("Over-fed" if c>tgt+1 else "In line")),al=left)
    r+=1
r+=1
head(ws,["Routing check (developers → Revati, channel partners → Ashwini)","Count","","Verdict"],r);r+=1
rt=S["routing"]
for lab,k,good in [("Developers → Revati (correct)","dev_to_revati",True),("Developers → Ashwini (mis-routed)","dev_to_ashwini",False),
                   ("Channel partners → Ashwini (correct)","cp_to_ashwini",True),("Channel partners → Revati (mis-routed)","cp_to_revati",False)]:
    cv(ws,r,1,lab,True,left);cv(ws,r,2,rt[k],True,fill=(GREENBG if good else REDBG),fc=(GREENF if good else REDF))
    cv(ws,r,3,"");cv(ws,r,4,("Good" if good else "Mis-routed"),al=left);r+=1
cv(ws,r+1,1,"Revati is carrying 17 channel-partner demos (her weak segment, 4.8% historical win) while Ashwini holds 8 developers (her weak segment, 3.7%). Routing is still inverted.",al=wrap)
ws.merge_cells(start_row=r+1,start_column=1,end_row=r+1,end_column=4);ws.row_dimensions[r+1].height=30

# 4 HOT list
ws=wb.create_sheet("HOT list");ws.sheet_view.showGridLines=False
cols=[("Prob",7),("Band",7),("Deal ID",10),("Name",20),("Seats",6),("City",14),("Type",13),("Owner",16),("Follow-ups",9),("Next action",62)]
for i,(h,w) in enumerate(cols):ws.column_dimensions[get_column_letter(i+1)].width=w
title(ws,"Chase this week — HOT + WARM",len(cols),TEAL)
head(ws,[c[0] for c in cols],2)
r=3
for d in sorted([x for x in rows if x["band"] in("HOT","WARM")],key=lambda x:-x["prob"]):
    vals=[d["prob"],d["band"],d["deal_id"],d["name"],d["licences"],d.get("city"),d.get("dev_broker"),d.get("owner"),d.get("sales_followup_since_demo"),d.get("next_action")]
    for ci,v in enumerate(vals,1):
        cv(ws,r,ci,v,al=(wrap if ci==10 else (left if ci in(4,6,7,8) else ctr)),nf=("0%" if ci==1 else None))
    ws.cell(row=r,column=2).fill=PatternFill("solid",fgColor=GREENBG if d["band"]=="HOT" else AMB)
    if (d.get("sales_followup_since_demo") or 0)==0:
        ws.cell(row=r,column=9).fill=PatternFill("solid",fgColor=REDBG);ws.cell(row=r,column=9).font=f(9,REDF,True)
    r+=1

# 5 All demos
ws=wb.create_sheet("All demos (43)");ws.freeze_panes="A3"
cols=[("Deal ID",10),("Name",20),("Creator",14),("Owner",16),("Source",10),("Seats",6),("Tier",7),("City",14),("Type",13),("Mode",8),("BANT",8),("Verdict",8),("Prob",7),("Band",7),("Follow-ups",9),("One-line",56)]
for i,(h,w) in enumerate(cols):ws.column_dimensions[get_column_letter(i+1)].width=w
title(ws,"All demos conducted 6–12 Jul 2026",len(cols))
head(ws,[c[0] for c in cols],2)
r=3
for d in sorted(rows,key=lambda x:-x["prob"]):
    vals=[d["deal_id"],d["name"],d["creator"],d.get("owner"),d["source"],d["licences"],d.get("tier"),d.get("city"),d.get("dev_broker"),d.get("onsite"),d.get("bant_flag"),d["verdict"],d["prob"],d["band"],d.get("sales_followup_since_demo"),d.get("one_line")]
    for ci,v in enumerate(vals,1):
        cv(ws,r,ci,v,al=(wrap if ci==16 else (left if ci in(2,3,4,8,9) else ctr)),nf=("0%" if ci==13 else None))
    vc=ws.cell(row=r,column=12);vd=d["verdict"]
    vc.fill=PatternFill("solid",fgColor={"GREEN":GREENBG,"AMBER":AMB,"RED":REDBG}[vd]);vc.font=f(9,{"GREEN":GREENF,"AMBER":"9C6500","RED":REDF}[vd],True)
    r+=1
ws.auto_filter.ref=f"A2:P{r-1}"
wb.save(f"{BASE}/Marketing_Weekly_{WK}.xlsx")
print("saved",wb.sheetnames)
