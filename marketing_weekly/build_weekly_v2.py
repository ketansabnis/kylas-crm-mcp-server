import json, glob, os, collections, statistics
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
BASE="/sessions/gifted-happy-euler/mnt/kylas-crm-mcp-server/marketing_weekly"
WK="2026-07-12"
rows=json.load(open(f"{BASE}/deals_{WK}.json"))
S=json.load(open(f"{BASE}/snapshot_{WK}.json"))
lost=[json.load(open(f)) for f in glob.glob(f"{BASE}/lost_{WK}/*.json") if os.path.basename(f)[:-5].isdigit()]
N=len(rows)
# --- new metrics ---
FUNNEL={"created":51,"demoed":34,"leak":17,"demos_conducted":43,"older_demoed":9}
CAMP={"Meta Lead-Gen (Apr-24)":4,"Meta Lead-Gen (Jun-25)":6,"Google Brand Search":3,"Google Search (jan24)":3,"Other/untagged paid":2}
COHORT={"demos":107,"lost":43,"won":2,"open":62}
# pre-sales rep scorecard
reps={}
for cr in set(r["creator"] for r in rows):
    sub=[r for r in rows if r["creator"]==cr]; n=len(sub)
    reps[cr]={"demos":n,
      "GREEN":sum(1 for r in sub if r["verdict"]=="GREEN"),"RED":sum(1 for r in sub if r["verdict"]=="RED"),
      "five":sum(1 for r in sub if r["is5plus"]),"dev":sum(1 for r in sub if r.get("dev_broker")=="Developer"),
      "onsite":sum(1 for r in sub if r.get("onsite")=="Onsite"),
      "hotwarm":sum(1 for r in sub if r["band"] in("HOT","WARM")),
      "badname":sum(1 for r in sub if r.get("name_valid") is False),
      "unver":sum(1 for r in sub if r.get("unverified_contact"))}
ent=[r for r in rows if (r["licences"] or 0)>=40]
# losses
def gi(d,k,dv=0):
    v=d.get(k)
    try: return int(v)
    except: return dv
lr=collections.Counter(str(d.get("loss_reason")) for d in lost)
comp=collections.Counter(str(d.get("competitor")) for d in lost if d.get("competitor") and str(d.get("competitor")).lower()!="none")
att=[gi(d,"sales_attempts") for d in lost]
zero=[d for d in lost if gi(d,"sales_attempts")==0]
S.update({"funnel":FUNNEL,"campaign":CAMP,"cohort":COHORT,"reps":reps,"enterprise40":len(ent),
          "losses":{"n":len(lost),"avg_attempts":round(statistics.mean(att),1),"zero_followup":len(zero),
                    "reasons":dict(lr),"competitors":dict(comp)}})
json.dump(S,open(f"{BASE}/snapshot_{WK}.json","w"),indent=2)

NAVY="1F3864";BLUE="2E5496";TEAL="1F6E6E";GREENBG="E2EFDA";AMB="FFF2CC";REDBG="FBE4E4";WHITE="FFFFFF";GREENF="1E7145";REDF="9C0006";GOLD="FFF4D6"
thin=Side(style="thin",color="BFBFBF");bd=Border(left=thin,right=thin,top=thin,bottom=thin)
def hf(s=10,c=WHITE,b=True):return Font(name="Arial",size=s,bold=b,color=c)
def f(s=9,c="000000",b=False):return Font(name="Arial",size=s,bold=b,color=c)
wrap=Alignment(wrap_text=True,vertical="top");ctr=Alignment(horizontal="center",vertical="center");left=Alignment(horizontal="left",vertical="center");ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=Workbook()
def title(ws,t,span,c=NAVY):
    x=ws.cell(row=1,column=1,value=t);x.font=hf(12,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=left
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=span);ws.row_dimensions[1].height=24
def head(ws,hs,row,c=BLUE):
    for i,h in enumerate(hs):
        x=ws.cell(row=row,column=i+1,value=h);x.font=hf(9,WHITE,True);x.fill=PatternFill("solid",fgColor=c);x.alignment=ctrw;x.border=bd
def cv(ws,r,c,v,b=False,al=None,fill=None,fc="000000",nf=None):
    x=ws.cell(row=r,column=c,value=v);x.border=bd;x.font=f(9,fc,b);x.alignment=al or ctr
    if fill:x.fill=PatternFill("solid",fgColor=fill)
    if nf:x.number_format=nf
T=S["targets"]
# 1 Scorecard
ws=wb.active;ws.title="Scorecard";ws.sheet_view.showGridLines=False
for c,w in zip("ABCDE",[34,13,12,10,54]):ws.column_dimensions[c].width=w
title(ws,f"Marketing & Pre-Sales weekly audit — {S['week_start']} to {S['week_end']}",5)
head(ws,["Metric","Actual","Target","Status","Read"],2)
def sc(r,m,a,t,ok,read):
    cv(ws,r,1,m,True,left);cv(ws,r,2,a,True,fill=(GREENBG if ok else REDBG),fc=(GREENF if ok else REDF))
    cv(ws,r,3,t);cv(ws,r,4,"ON" if ok else "MISS",True,fc=(GREENF if ok else REDF));cv(ws,r,5,read,al=wrap)
r=3
sc(r,"Deals created",FUNNEL["created"],"—",True,"Top of funnel volume");r+=1
sc(r,"Demo actually conducted (of those)",f'{FUNNEL["demoed"]} ({FUNNEL["demoed"]/FUNNEL["created"]:.0%})',"≥85%",False,"17 created but NO demo — a third of new leads leaked in one week");r+=1
sc(r,"Demos conducted (all)",N,"—",True,"43 (34 new + 9 older). Down vs ~59/wk June run-rate");r+=1
sc(r,"Demos of 5+ licences",S["five"],T["five"],S["five"]>=T["five"],"52% of target");r+=1
sc(r,"GREEN (well-qualified)",S["green"],T["green"],False,"16% of demos (June: 39%)");r+=1
sc(r,"RED (shouldn't be booked)",S["verdict"].get("RED",0),"<10%",False,"37% RED — worse than June's 22%");r+=1
sc(r,"Developer share",f'{S["dev"]}/{N} ({S["dev"]/N:.0%})',"≥50%",False,"Below floor — CP-heavy");r+=1
sc(r,"Onsite (Main-6)",f'{S["onsite_m6"]}/{S["main6"]}',"≥40%",False,"15% — on-ground model still not happening");r+=1
sc(r,"Enterprise (40+ seats)",S["enterprise40"],"≥1",S["enterprise40"]>=1,"Only 1 enterprise-scale demo — enterprise pipeline still near-absent");r+=1
sc(r,"HOT + WARM",S["band"].get("HOT",0)+S["band"].get("WARM",0),"—",True,"7 worth chasing; 84% landed COLD");r+=1
r+=1
cv(ws,r,1,"COHORT ROT — prior fortnight's demos (22 Jun–5 Jul)",True,left,fill=REDBG);cv(ws,r,2,"");cv(ws,r,3,"");cv(ws,r,4,"");cv(ws,r,5,"");r+=1
for lab,v,rd in [("Demos in that cohort",COHORT["demos"],""),
                 ("Already lost / unqualified",f'{COHORT["lost"]} ({COHORT["lost"]/COHORT["demos"]:.0%})',"40% dead within 1–3 weeks of the demo"),
                 ("Won",COHORT["won"],"2 wins from 107 demos"),
                 ("Still open",COHORT["open"],"")]:
    cv(ws,r,1,lab,True,left);cv(ws,r,2,v,True,fill=(REDBG if "lost" in lab.lower() else None));cv(ws,r,3,"");cv(ws,r,4,"");cv(ws,r,5,rd,al=wrap);r+=1

# 2 Pre-sales rep scorecard
ws=wb.create_sheet("Pre-Sales Rep Scorecard");ws.sheet_view.showGridLines=False
for c,w in zip("ABCDEFGHI",[18,8,8,8,9,8,9,10,10]):ws.column_dimensions[c].width=w
title(ws,"Pre-sales rep quality — who is generating what",9,TEAL)
head(ws,["Rep","Demos","GREEN","RED","%GREEN","5+","Developer","Onsite","HOT/WARM"],2)
r=3
for cr,v in sorted(reps.items(),key=lambda kv:-kv[1]["demos"]):
    cv(ws,r,1,cr,True,left);cv(ws,r,2,v["demos"]);cv(ws,r,3,v["GREEN"]);cv(ws,r,4,v["RED"],fill=REDBG if v["RED"]/v["demos"]>0.3 else None)
    cv(ws,r,5,v["GREEN"]/v["demos"],nf="0%");cv(ws,r,6,v["five"]);cv(ws,r,7,v["dev"]);cv(ws,r,8,v["onsite"]);cv(ws,r,9,v["hotwarm"]);r+=1
r+=1
head(ws,["Data hygiene by rep","Junk names","Unverified contact","","","","","",""],r);r+=1
for cr,v in sorted(reps.items(),key=lambda kv:-kv[1]["demos"]):
    cv(ws,r,1,cr,True,left);cv(ws,r,2,v["badname"],fill=REDBG if v["badname"] else None);cv(ws,r,3,v["unver"],fill=REDBG if v["unver"] else None)
    for c in range(4,10): cv(ws,r,c,"")
    r+=1

# 3 Funnel & Campaigns
ws=wb.create_sheet("Funnel & Campaigns");ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[34,12,12,56]):ws.column_dimensions[c].width=w
title(ws,"Top-of-funnel: created → demo, and paid campaigns",4,TEAL)
head(ws,["Funnel step","Count","%","Read"],2)
r=3
cv(ws,r,1,"Deals created by pre-sales",True,left);cv(ws,r,2,FUNNEL["created"]);cv(ws,r,3,"100%");cv(ws,r,4,"",al=wrap);r+=1
cv(ws,r,1,"— of which demo conducted",True,left);cv(ws,r,2,FUNNEL["demoed"],fill=AMB);cv(ws,r,3,FUNNEL["demoed"]/FUNNEL["created"],nf="0%");cv(ws,r,4,"Show-rate 67%",al=wrap);r+=1
cv(ws,r,1,"— LEAK: created, no demo",True,left);cv(ws,r,2,FUNNEL["leak"],fill=REDBG,fc=REDF);cv(ws,r,3,FUNNEL["leak"]/FUNNEL["created"],nf="0%");cv(ws,r,4,"A third of new leads never reached a demo — invisible until now",al=wrap);r+=1
r+=1
head(ws,["Paid campaign","Demos","","Read"],r);r+=1
cn={"Meta Lead-Gen (Apr-24)":"Meta lead forms — biggest paid cluster, and Facebook produced ZERO HOT/WARM this week",
    "Meta Lead-Gen (Jun-25)":"Meta lead forms","Google Brand Search":"High-intent, keep","Google Search (jan24)":"Prospecting search","Other/untagged paid":"Check tagging"}
for k,v in CAMP.items():
    fill=REDBG if k.startswith("Meta") else (GREENBG if k.startswith("Google") else None)
    cv(ws,r,1,k,True,left,fill);cv(ws,r,2,v);cv(ws,r,3,"");cv(ws,r,4,cn.get(k,""),al=wrap);r+=1
cv(ws,r,1,"Meta total",True,left,REDBG);cv(ws,r,2,10,True,fill=REDBG,fc=REDF);cv(ws,r,3,"");cv(ws,r,4,"10 of 18 tagged paid demos are Meta — and Meta/Facebook delivered 1 GREEN, 5 RED, 0 HOT/WARM",al=wrap);r+=1
cv(ws,r,1,"Google total",True,left,GREENBG);cv(ws,r,2,6,True,fill=GREENBG,fc=GREENF);cv(ws,r,3,"");cv(ws,r,4,"Best size mix (7 of 10 Google-source demos were 5+)",al=wrap)

# 4 Source
ws=wb.create_sheet("Source");ws.sheet_view.showGridLines=False
for c,w in zip("ABCDEFGH",[14,9,9,9,9,9,11,46]):ws.column_dimensions[c].width=w
title(ws,"Channel performance",8,TEAL)
head(ws,["Source","Demos","5+","GREEN","RED","Dev","HOT/WARM","Read"],2)
notes={"Organic":"Best channel — most 5+, most developers, 4 of 7 HOT/WARM. Protect & grow.",
 "Google":"Best size mix (7 of 10 are 5+). Scale.",
 "Facebook":"11 demos → 1 GREEN, 5 RED, 1 developer, ZERO HOT/WARM. Third confirmation. Cut or re-target.",
 "WhatsApp":"Only 1 of 7 is 5+, 4 RED — inbound WA bringing tiny/unqualified leads. New watch-list item.",
 "Unknown":"Source not set — fix tagging."}
r=3
for s,v in sorted(S["source"].items(),key=lambda kv:-kv[1]["demos"]):
    fill=REDBG if s in("Facebook","WhatsApp") else (GREENBG if s in("Organic","Google") else None)
    cv(ws,r,1,s,True,left,fill);cv(ws,r,2,v["demos"]);cv(ws,r,3,v["five"]);cv(ws,r,4,v["green"]);cv(ws,r,5,v["red"]);cv(ws,r,6,v["dev"]);cv(ws,r,7,v["hotwarm"]);cv(ws,r,8,notes.get(s,""),al=wrap);r+=1

# 5 Losses
ws=wb.create_sheet("Losses");ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[28,10,14,56]):ws.column_dimensions[c].width=w
title(ws,f"Deals lost last week ({len(lost)}, 5+ seats) — why, and how hard sales chased",4,TEAL)
head(ws,["Loss reason","Deals","Avg attempts","Read"],2)
r=3
for k,v in lr.most_common():
    sub=[gi(d,"sales_attempts") for d in lost if str(d.get("loss_reason"))==k]
    cv(ws,r,1,k,True,left);cv(ws,r,2,v);cv(ws,r,3,round(statistics.mean(sub),1) if sub else "")
    cv(ws,r,4,("Price is the real theme — merge with Budget" if k.startswith("Competitor") else ""),al=wrap);r+=1
r+=1
head(ws,["Competitor we lost to","Deals","","Read"],r);r+=1
for k,v in comp.most_common(8):
    cv(ws,r,1,k,True,left,REDBG if v>1 else None);cv(ws,r,2,v);cv(ws,r,3,"")
    cv(ws,r,4,("PRICE UNDERCUT — Rs300–500/user vs our Rs1000–1700" if k.lower()=="buildesk" else ""),al=wrap);r+=1
r+=1
cv(ws,r,1,"Avg sales attempts before giving up",True,left);cv(ws,r,2,round(statistics.mean(att),1),True,fill=GOLD);cv(ws,r,3,"");cv(ws,r,4,"Bimodal: some deals chased 20–31 times, others killed with zero follow-up",al=wrap);r+=1
cv(ws,r,1,"Lost with ZERO follow-up",True,left);cv(ws,r,2,len(zero),True,fill=REDBG,fc=REDF);cv(ws,r,3,"");cv(ws,r,4,"Includes a 35-licence Bangalore deal that went to Salesforce with no post-demo contact",al=wrap);r+=1
r+=2
head(ws,["Deal","Owner","Attempts","Why lost"],r);r+=1
for d in sorted(lost,key=lambda x:-gi(x,"sales_attempts")):
    cv(ws,r,1,f'#{d["deal_id"]} {str(d.get("name"))[:20]}',al=left);cv(ws,r,2,str(d.get("owner"))[:18],al=left)
    a=gi(d,"sales_attempts");cv(ws,r,3,a,fill=REDBG if a==0 else None,fc=REDF if a==0 else "000000")
    cv(ws,r,4,d.get("one_line"),al=wrap);r+=1

# 6 HOT list
ws=wb.create_sheet("HOT list");ws.sheet_view.showGridLines=False
cols=[("Prob",7),("Band",7),("Deal ID",10),("Name",20),("Seats",6),("City",14),("Type",13),("Owner",16),("Follow-ups",9),("Next action",62)]
for i,(h,w) in enumerate(cols):ws.column_dimensions[get_column_letter(i+1)].width=w
title(ws,"Chase this week — HOT + WARM",len(cols),TEAL)
head(ws,[c[0] for c in cols],2)
r=3
for d in sorted([x for x in rows if x["band"] in("HOT","WARM")],key=lambda x:-x["prob"]):
    vals=[d["prob"],d["band"],d["deal_id"],d["name"],d["licences"],d.get("city"),d.get("dev_broker"),d.get("owner"),d.get("sales_followup_since_demo"),d.get("next_action")]
    for ci,v in enumerate(vals,1):
        cv(ws,r,ci,v,al=(wrap if ci==10 else (left if ci in(4,6,7,8) else ctr)),nf=("0%" if ci==1 else None))
    ws.cell(row=r,column=2).fill=PatternFill("solid",fgColor=GREENBG if d["band"]=="HOT" else AMB)
    if (d.get("sales_followup_since_demo") or 0)==0:
        ws.cell(row=r,column=9).fill=PatternFill("solid",fgColor=REDBG);ws.cell(row=r,column=9).font=f(9,REDF,True)
    r+=1

# 7 Allocation & Routing
ws=wb.create_sheet("Allocation & Routing");ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[26,12,14,52]):ws.column_dimensions[c].width=w
title(ws,"Sales allocation, and pre-sales lead routing",4,TEAL)
head(ws,["Sales owner","Demos","Even-split","Note"],2)
tgt=round(N/len(S["alloc"]),1);r=3
for o,c in sorted(S["alloc"].items(),key=lambda kv:-kv[1]):
    cv(ws,r,1,o,True,left);cv(ws,r,2,c);cv(ws,r,3,tgt);cv(ws,r,4,("Under-fed" if c<tgt-1 else ("Over-fed" if c>tgt+1 else "In line")),al=left);r+=1
r+=1
head(ws,["Routing (developers→Revati, channel partners→Ashwini)","Count","","Verdict"],r);r+=1
rt=S["routing"]
for lab,k,good in [("Developers → Revati (correct)","dev_to_revati",True),("Developers → Ashwini (mis-routed)","dev_to_ashwini",False),
                   ("Channel partners → Ashwini (correct)","cp_to_ashwini",True),("Channel partners → Revati (mis-routed)","cp_to_revati",False)]:
    cv(ws,r,1,lab,True,left);cv(ws,r,2,rt[k],True,fill=(GREENBG if good else REDBG),fc=(GREENF if good else REDF));cv(ws,r,3,"");cv(ws,r,4,("Good" if good else "Mis-routed"),al=left);r+=1

# 8 All demos
ws=wb.create_sheet("All demos (43)");ws.freeze_panes="A3"
cols=[("Deal ID",10),("Name",20),("Creator",14),("Owner",16),("Source",10),("Seats",6),("Tier",7),("City",14),("Type",13),("Mode",8),("BANT",8),("Verdict",8),("Prob",7),("Band",7),("Follow-ups",9),("One-line",56)]
for i,(h,w) in enumerate(cols):ws.column_dimensions[get_column_letter(i+1)].width=w
title(ws,"All demos conducted 6–12 Jul 2026",len(cols))
head(ws,[c[0] for c in cols],2)
r=3
for d in sorted(rows,key=lambda x:-x["prob"]):
    vals=[d["deal_id"],d["name"],d["creator"],d.get("owner"),d["source"],d["licences"],d.get("tier"),d.get("city"),d.get("dev_broker"),d.get("onsite"),d.get("bant_flag"),d["verdict"],d["prob"],d["band"],d.get("sales_followup_since_demo"),d.get("one_line")]
    for ci,v in enumerate(vals,1):
        cv(ws,r,ci,v,al=(wrap if ci==16 else (left if ci in(2,3,4,8,9) else ctr)),nf=("0%" if ci==13 else None))
    vc=ws.cell(row=r,column=12);vd=d["verdict"]
    vc.fill=PatternFill("solid",fgColor={"GREEN":GREENBG,"AMBER":AMB,"RED":REDBG}[vd]);vc.font=f(9,{"GREEN":GREENF,"AMBER":"9C6500","RED":REDF}[vd],True)
    r+=1
ws.auto_filter.ref=f"A2:P{r-1}"
wb.save(f"{BASE}/Marketing_Weekly_{WK}.xlsx")
print("saved:",wb.sheetnames)
print("reps:",json.dumps(reps,indent=1))
print("enterprise40:",len(ent))
print("losses:",len(lost),"avg attempts",round(statistics.mean(att),1),"zero-fu",len(zero))
print("reasons:",dict(lr))
print("competitors:",dict(comp))
